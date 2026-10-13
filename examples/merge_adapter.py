"""基础 XLSX 字段合并示例，不包含订单去重或电商对账规则。"""

import argparse
import json
from pathlib import Path
import shutil
import sys
import tempfile

from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell


EXCEL_MAX_ROWS = 1_048_576
SOURCE_COLUMNS = ("_来源文件", "_来源工作表", "_来源行号")


def _empty(value):
    return value is None or (isinstance(value, str) and value == "")


def _cells(worksheet, values):
    for value in values:
        if isinstance(value, str):
            cell = WriteOnlyCell(worksheet, value=value)
            # A genuine text value starting with '=' must remain text.
            cell.data_type = "s"
            yield cell
        else:
            yield value


class _SplitWriter:
    def __init__(self, directory, header, max_rows):
        self.directory = directory
        self.header = header
        self.max_rows = max_rows
        self.outputs = []
        self.workbook = None
        self.worksheet = None
        self.rows = 0
        self._new_part()

    def _new_part(self):
        self.workbook = Workbook(write_only=True)
        self.worksheet = self.workbook.create_sheet("合并结果")
        self.worksheet.append(_cells(self.worksheet, self.header))
        self.rows = 1

    def _save(self):
        name = f"merged-{len(self.outputs) + 1:03d}.xlsx"
        self.workbook.save(self.directory / name)
        self.workbook.close()
        self.workbook = None
        self.outputs.append(name)

    def append(self, values):
        if self.rows >= self.max_rows:
            self._save()
            self._new_part()
        self.worksheet.append(_cells(self.worksheet, values))
        self.rows += 1

    def finish(self):
        if self.workbook is not None:
            self._save()

    def discard(self):
        # Saving only inside staging also closes openpyxl's streaming writers.
        if self.workbook is not None:
            try:
                self.workbook.save(self.directory / "_discarded.xlsx")
            finally:
                self.workbook.close()
                self.workbook = None


def merge_files(manifest, output_dir, *, max_rows=EXCEL_MAX_ROWS, on_progress=None):
    """逐行合并，全部输入通过验证后一次发布整个成果目录。"""
    files = manifest.get("files") if isinstance(manifest, dict) else None
    if not isinstance(files, list) or not files:
        raise ValueError("文件清单必须包含至少一份 XLSX 文件")
    if not isinstance(max_rows, int) or not 2 <= max_rows <= EXCEL_MAX_ROWS:
        raise ValueError("输出行数上限必须在 2 至 1,048,576 之间，包含表头")
    output_dir = Path(output_dir).resolve()
    output_dir.parent.mkdir(parents=True, exist_ok=True)
    if output_dir.exists() and (not output_dir.is_dir() or any(output_dir.iterdir())):
        raise ValueError("输出目录必须为空，请为每次尝试使用新的目录")
    staging = Path(tempfile.mkdtemp(prefix=".merge-", dir=output_dir.parent))
    writer = None
    expected_keys = None
    header = None
    sheet_count = 0
    record_count = 0

    def progress(phase, current, message):
        if on_progress is not None:
            on_progress({"type": "progress", "phase": phase, "current": current,
                         "total": len(files), "message": message})

    try:
        for file_index, entry in enumerate(files, start=1):
            if not isinstance(entry, dict) or not isinstance(entry.get("path"), str):
                raise ValueError(f"清单第 {file_index} 个文件缺少有效的 path")
            source = entry.get("relativePath") or entry.get("name") or Path(entry["path"]).name
            if not isinstance(source, str):
                raise ValueError(f"清单第 {file_index} 个文件的来源名称无效")
            progress("parsing", file_index - 1, f"正在读取：{source}")
            workbook = load_workbook(entry["path"], read_only=True, data_only=False)
            try:
                for worksheet in workbook.worksheets:
                    # Read actual rows rather than trusting an export's dimension hint.
                    worksheet.reset_dimensions()
                    indexes = None
                    column_count = 0
                    for row_number, row in enumerate(worksheet.iter_rows(), start=1):
                        location = f"{source} / {worksheet.title} 第 {row_number} 行"
                        if any(cell.data_type == "f" for cell in row):
                            raise ValueError(f"发现公式单元格：{location}；基础适配器不计算公式，请提供值数据")
                        values = [cell.value for cell in row]
                        if all(_empty(value) for value in values):
                            continue
                        if indexes is None:
                            while values and _empty(values[-1]):
                                values.pop()
                            if any(_empty(value) or (isinstance(value, str) and not value.strip())
                                   for value in values):
                                raise ValueError(f"表头存在空字段：{location}")
                            keys = [(type(value).__name__, value) for value in values]
                            if len(set(keys)) != len(keys):
                                raise ValueError(f"表头存在重复字段：{location}")
                            if any(value in SOURCE_COLUMNS for value in values):
                                raise ValueError(f"表头与保留来源字段重名：{location}")
                            if expected_keys is None:
                                expected_keys, header = keys, values
                                writer = _SplitWriter(staging, header + list(SOURCE_COLUMNS), max_rows)
                            elif set(keys) != set(expected_keys):
                                missing = [repr(value) for key_type, value in expected_keys if (key_type, value) not in set(keys)]
                                extra = [repr(value) for key_type, value in keys if (key_type, value) not in set(expected_keys)]
                                raise ValueError(f"字段不一致：{location}；缺少 {', '.join(missing) or '无'}；新增 {', '.join(extra) or '无'}")
                            positions = {key: index for index, key in enumerate(keys)}
                            indexes = [positions[key] for key in expected_keys]
                            column_count = len(values)
                            sheet_count += 1
                            continue
                        if any(not _empty(value) for value in values[column_count:]):
                            raise ValueError(f"存在没有对应表头的数据列：{location}")
                        values.extend([None] * max(0, column_count - len(values)))
                        writer.append([values[index] for index in indexes] + [source, worksheet.title, row_number])
                        record_count += 1
            finally:
                workbook.close()
            progress("parsing", file_index, f"已读取：{source}")

        if writer is None:
            raise ValueError("所有工作表均为空，至少需要一张含有效表头的数据表")
        progress("exporting", len(files), "输入检查完成，正在导出统一结果表")
        writer.finish()
        summary = {
            "fileCount": len(files), "sheetCount": sheet_count, "recordCount": record_count,
            "outputCount": len(writer.outputs), "description": "基础字段合并；无业务去重",
        }
        (staging / "summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
        # The executor creates an empty output directory; replace it with one
        # complete staged directory rather than exposing individual partial files.
        if output_dir.exists():
            output_dir.rmdir()
        staging.rename(output_dir)
        staging = None
        return summary
    finally:
        if staging is not None:
            try:
                if writer is not None:
                    writer.discard()
            finally:
                shutil.rmtree(staging)


def main():
    for stream in (sys.stdout, sys.stderr):
        if hasattr(stream, "reconfigure"):
            stream.reconfigure(encoding="utf-8")
    parser = argparse.ArgumentParser(description="基础 XLSX 字段合并适配器，不做业务去重")
    parser.add_argument("--manifest", required=True)
    parser.add_argument("--output-dir", required=True)
    args = parser.parse_args()
    try:
        manifest = json.loads(Path(args.manifest).read_text(encoding="utf-8"))
        merge = merge_files
        if manifest.get("workflow"):
            from merge_workflow import merge_prepared
            merge = merge_prepared
        merge(manifest, args.output_dir,
                    on_progress=lambda value: print(json.dumps(value, ensure_ascii=False), flush=True))
    except Exception as error:
        print(f"合并失败：{error}", file=sys.stderr, flush=True)
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
