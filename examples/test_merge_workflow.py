import hashlib
import json
from pathlib import Path
import tempfile
import unittest

from openpyxl import Workbook, load_workbook

from merge_workflow import merge_prepared, prepare_files, preview_workbook, validate_outputs


class WorkflowTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory(prefix="pi-workflow-")
        self.root = Path(self.temporary.name)

    def tearDown(self):
        self.temporary.cleanup()

    def workbook(self, name, sheets):
        path = self.root / name
        workbook = Workbook()
        workbook.remove(workbook.active)
        for title, rows in sheets.items():
            worksheet = workbook.create_sheet(title)
            for row in rows:
                worksheet.append(row)
        workbook.save(path)
        workbook.close()
        return path

    def manifest(self, *paths, rules=None):
        return {"batchId": "test", "files": [{"id": str(index), "name": path.name, "relativePath": f"shop/{path.name}", "path": str(path), "sha256": hashlib.sha256(path.read_bytes()).hexdigest()} for index, path in enumerate(paths)], "excluded": [],
                "workflow": {"cacheDir": str(self.root / "checkpoints"), "inputRevision": 1, "rules": rules or {"columns": [], "mappings": [], "sheetNames": []}}}

    def test_confirmed_mapping_and_sheet_selection_preserve_strings_and_sources(self):
        source = self.workbook("a.xlsx", {"订单": [["交易号", "支付金额", "附注"], ["000123", 2.5, "x"]], "统计": [["说明", None, "无效表头"], ["跳过"]]})
        rules = {"columns": ["订单号", "金额"], "mappings": [{"headers": ["交易号", "支付金额", "附注"], "targets": {"交易号": "订单号", "支付金额": "金额", "附注": None}}], "sheetNames": ["订单"]}
        manifest = self.manifest(source, rules=rules)
        report = prepare_files(manifest)
        self.assertTrue(report["ok"], report)
        self.assertEqual(report["recordCount"], 1)
        self.assertEqual({sheet["name"] for sheet in report["files"][0]["sheets"]}, {"订单", "统计"})
        output = self.root / "output"
        output.mkdir()
        summary = merge_prepared(manifest, output)
        self.assertEqual(summary["cachedCount"], 1)
        workbook = load_workbook(output / "merged-001.xlsx", read_only=True)
        try:
            rows = list(workbook.active.values)
            self.assertEqual(rows[1], ("000123", 2.5, "shop/a.xlsx", "订单", 2))
        finally:
            workbook.close()

    def test_collects_all_failed_files_and_retries_only_affected_parsing(self):
        good = self.workbook("good.xlsx", {"订单": [["ID"], ["0001"]]})
        formula = self.workbook("formula.xlsx", {"订单": [["ID"], ["=1+1"]]})
        mismatch = self.workbook("mismatch.xlsx", {"订单": [["Other"], ["0003"]]})
        manifest = self.manifest(good, formula, mismatch)
        first = prepare_files(manifest)
        self.assertFalse(first["ok"])
        self.assertEqual(first["successCount"], 1)
        self.assertEqual(first["failedCount"], 2)
        self.assertEqual(first["issues"][0]["row"], 2)
        self.assertEqual(first["issues"][0]["sheet"], "订单")
        checkpoint = self.root / "checkpoints" / first["files"][0]["cacheKey"] / "metadata.json"
        modified = checkpoint.stat().st_mtime_ns
        self.workbook("formula.xlsx", {"订单": [["ID"], ["0002"]]})
        self.workbook("mismatch.xlsx", {"订单": [["ID"], ["0003"]]})
        for entry in manifest["files"]:
            entry["sha256"] = hashlib.sha256(Path(entry["path"]).read_bytes()).hexdigest()
        second = prepare_files(manifest)
        self.assertTrue(second["ok"], second)
        self.assertEqual(second["cachedCount"], 1)
        self.assertEqual(checkpoint.stat().st_mtime_ns, modified)
        output = self.root / "output"
        output.mkdir()
        merge_prepared(manifest, output)
        self.assertEqual(validate_outputs([output / "merged-001.xlsx"], 3)["recordCount"], 3)

    def test_rule_changes_and_checkpoint_corruption_force_rebuild(self):
        source = self.workbook("a.xlsx", {"订单": [["ID", "金额"], ["0001", 2]]})
        manifest = self.manifest(source)
        first = prepare_files(manifest)
        self.assertEqual(prepare_files(manifest)["cachedCount"], 1)
        cached = self.root / "checkpoints" / first["files"][0]["cacheKey"] / "merged-001.xlsx"
        cached.write_bytes(b"corrupt")
        self.assertEqual(prepare_files(manifest)["cachedCount"], 0)
        manifest["workflow"]["rules"]["columns"] = ["金额", "ID"]
        reordered = prepare_files(manifest)
        self.assertTrue(reordered["ok"])
        self.assertEqual(reordered["cachedCount"], 0)
        self.assertNotEqual(first["files"][0]["cacheKey"], reordered["files"][0]["cacheKey"])

    def test_changed_source_never_uses_old_upload_hash(self):
        source = self.workbook("a.xlsx", {"订单": [["ID"], ["0001"]]})
        manifest = self.manifest(source)
        self.assertTrue(prepare_files(manifest)["ok"])
        self.workbook("a.xlsx", {"订单": [["ID"], ["different"]]})
        report = prepare_files(manifest)
        self.assertFalse(report["ok"])
        self.assertIn("上传校验值", report["issues"][0]["message"])

    def test_preview_pagination_and_validation_reject_invalid_or_mismatched_results(self):
        source = self.workbook("result.xlsx", {"订单": [["ID"]] + [[f"00{index}"] for index in range(29)]})
        first = preview_workbook(source)
        self.assertEqual(len(first["rows"]), 25)
        self.assertEqual(first["rows"][0][0], "000")
        self.assertTrue(first["hasMore"])
        second = preview_workbook(source, offset=25)
        self.assertEqual(len(second["rows"]), 4)
        self.assertFalse(second["hasMore"])
        self.assertEqual(validate_outputs([source], 29)["recordCount"], 29)
        with self.assertRaisesRegex(ValueError, "记录数不一致"):
            validate_outputs([source], 30)
        with self.assertRaisesRegex(ValueError, "成果字段"):
            validate_outputs([source], 29, ["Wrong"])
        broken = self.root / "broken.xlsx"
        broken.write_text("not a workbook")
        with self.assertRaises(Exception):
            validate_outputs([broken])
        with self.assertRaisesRegex(ValueError, "工作表不存在"):
            preview_workbook(source, sheet="不存在")

    def test_duplicate_warning_exclusion_accounting_and_restart_cache(self):
        source = self.workbook("a.xlsx", {"订单": [["ID"], ["0001"]]})
        other = self.root / "b.xlsx"
        other.write_bytes(source.read_bytes())
        manifest = self.manifest(source, other)
        manifest["excluded"] = [{"id": "excluded", "relativePath": "old.xlsx", "reason": "不属于本期"}]
        report = prepare_files(manifest)
        self.assertEqual(len(report["duplicates"]), 1)
        self.assertEqual(report["recordCount"], 2)
        restarted = json.loads(json.dumps(manifest))
        self.assertEqual(prepare_files(restarted)["cachedCount"], 2)
        output = self.root / "output"
        output.mkdir()
        summary = merge_prepared(restarted, output)
        self.assertEqual(summary["excluded"][0]["reason"], "不属于本期")


if __name__ == "__main__":
    unittest.main()
