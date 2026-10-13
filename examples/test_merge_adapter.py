import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest

from openpyxl import Workbook, load_workbook

from merge_adapter import EXCEL_MAX_ROWS, SOURCE_COLUMNS, merge_files


class Text(str):
    """Test fixture for real string cells whose values look like formulae."""


class MergeAdapterTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory(prefix="pi-merge-test-")
        self.root = Path(self.temporary.name)

    def tearDown(self):
        self.temporary.cleanup()

    def workbook(self, name, sheets):
        path = self.root / name
        workbook = Workbook()
        workbook.remove(workbook.active)
        for title, rows in sheets.items():
            worksheet = workbook.create_sheet(title)
            for row_number, values in enumerate(rows, start=1):
                for column_number, value in enumerate(values, start=1):
                    cell = worksheet.cell(row_number, column_number, value=value)
                    if isinstance(value, Text):
                        cell.data_type = "s"
        workbook.save(path)
        workbook.close()
        return path

    def manifest(self, *paths):
        return {"batchId": "test-batch", "files": [
            {"name": path.name, "relativePath": f"店铺/{path.name}", "path": str(path)}
            for path in paths
        ]}

    def result(self, output, name="merged-001.xlsx"):
        workbook = load_workbook(output / name, read_only=True, data_only=False)
        try:
            return list(workbook.active.iter_rows(values_only=True))
        finally:
            workbook.close()

    def assert_no_results(self, output):
        self.assertEqual(list(output.glob("*.xlsx")), [])
        self.assertEqual(list(self.root.glob(".merge-*")), [])

    def test_reordered_fields_strings_empty_rows_and_sources(self):
        long_text = "中文长文本" * 3000
        first = self.workbook("中文订单.xlsx", {
            "订单": [
                [None, None, None, None],
                ["订单号", "商品", "金额", "备注"],
                ["00123", "水杯", 10, long_text],
                [None, None, None, None],
                ["00000", "书", 0, Text("=真实文字")],
            ],
            "空工作表": [],
        })
        second = self.workbook("另一店铺.xlsx", {
            "订单明细": [
                ["金额", "订单号", "备注", "商品"],
                [2.5, "123456789012345678901234", None, "纸"],
            ],
        })
        events = []
        output = self.root / "output"
        summary = merge_files(self.manifest(first, second), output, on_progress=events.append)
        rows = self.result(output)
        self.assertEqual(rows[0], ("订单号", "商品", "金额", "备注") + SOURCE_COLUMNS)
        self.assertEqual(rows[1], ("00123", "水杯", 10, long_text, "店铺/中文订单.xlsx", "订单", 3))
        self.assertEqual(rows[2], ("00000", "书", 0, "=真实文字", "店铺/中文订单.xlsx", "订单", 5))
        self.assertEqual(rows[3], ("123456789012345678901234", "纸", 2.5, None, "店铺/另一店铺.xlsx", "订单明细", 2))
        self.assertEqual(summary, {
            "fileCount": 2, "sheetCount": 2, "recordCount": 3,
            "outputCount": 1, "description": "基础字段合并；无业务去重",
        })
        self.assertEqual(json.loads((output / "summary.json").read_text(encoding="utf-8")), summary)
        self.assertEqual(events[-1]["phase"], "exporting")
        self.assertEqual(events[-1]["current"], 2)
        workbook = load_workbook(output / "merged-001.xlsx", read_only=True, data_only=False)
        try:
            self.assertEqual(workbook.active["D3"].data_type, "s")
            self.assertEqual(workbook.active["D3"].value, "=真实文字")
        finally:
            workbook.close()

    def test_inconsistent_fields_fail_without_published_partial_results(self):
        first = self.workbook("first.xlsx", {"订单": [["订单号", "金额"], ["001", 2]]})
        second = self.workbook("second.xlsx", {"订单": [["订单号", "数量"], ["002", 3]]})
        output = self.root / "output"
        output.mkdir()
        with self.assertRaisesRegex(ValueError, "字段不一致.*second.xlsx.*缺少.*金额.*新增.*数量"):
            merge_files(self.manifest(first, second), output)
        self.assert_no_results(output)
        self.assertFalse((output / "summary.json").exists())

    def test_formula_cells_fail_and_are_never_evaluated(self):
        source = self.workbook("formula.xlsx", {
            "订单": [["订单号", "金额"], ["001", 2], ["002", "=SUM(1,2)"]],
        })
        output = self.root / "output"
        with self.assertRaisesRegex(ValueError, "发现公式单元格.*第 3 行.*不计算公式"):
            merge_files(self.manifest(source), output)
        self.assert_no_results(output)

    def test_invalid_headers_and_unheaded_data_fail(self):
        examples = [
            ([["订单号", "订单号"], ["001", "002"]], "重复字段"),
            ([["订单号", None, "金额"], ["001", None, 2]], "空字段"),
            ([["订单号", "_来源文件"], ["001", "x"]], "来源字段重名"),
            ([["订单号"], ["001", 2]], "没有对应表头"),
        ]
        for index, (rows, message) in enumerate(examples):
            with self.subTest(message=message):
                source = self.workbook(f"invalid-{index}.xlsx", {"订单": rows})
                output = self.root / f"output-{index}"
                with self.assertRaisesRegex(ValueError, message):
                    merge_files(self.manifest(source), output)
                self.assert_no_results(output)

    def test_split_respects_row_limit_including_header(self):
        source = self.workbook("split.xlsx", {
            "订单": [["订单号"]] + [[f"00{index}"] for index in range(5)],
        })
        output = self.root / "output"
        summary = merge_files(self.manifest(source), output, max_rows=3)
        self.assertEqual(summary["outputCount"], 3)
        outputs = sorted(output.glob("*.xlsx"))
        tables = [self.result(output, path.name) for path in outputs]
        self.assertEqual([len(table) for table in tables], [3, 3, 2])
        self.assertTrue(all(table[0] == ("订单号",) + SOURCE_COLUMNS for table in tables))
        self.assertEqual([row[0] for table in tables for row in table[1:]], [f"00{index}" for index in range(5)])
        self.assertEqual(EXCEL_MAX_ROWS, 1_048_576)

    def test_separate_attempts_do_not_accumulate_or_deduplicate_business_records(self):
        source = self.workbook("same-order.xlsx", {
            "订单": [["订单号", "金额"], ["001", 5], ["001", 5]],
        })
        manifest = self.manifest(source)
        first, second = self.root / "attempt-1", self.root / "attempt-2"
        summary_a = merge_files(manifest, first)
        summary_b = merge_files(manifest, second)
        self.assertEqual(summary_a, summary_b)
        self.assertEqual(summary_b["recordCount"], 2)
        self.assertEqual(self.result(first), self.result(second))
        self.assertEqual(len(self.result(second)), 3)

    def test_requires_at_least_one_table_and_accepts_header_only_table(self):
        empty = self.workbook("empty.xlsx", {"空": [[None, None]]})
        output = self.root / "output"
        with self.assertRaisesRegex(ValueError, "所有工作表均为空"):
            merge_files(self.manifest(empty), output)
        self.assert_no_results(output)
        header_only = self.workbook("header-only.xlsx", {"表": [["订单号"]]})
        summary = merge_files(self.manifest(empty, header_only), output)
        self.assertEqual(summary["sheetCount"], 1)
        self.assertEqual(summary["recordCount"], 0)
        self.assertEqual(len(self.result(output)), 1)

    def test_cli_reports_jsonl_progress_and_clear_chinese_errors(self):
        script = Path(__file__).with_name("merge_adapter.py")
        for index, rows in enumerate([[["订单号"], ["001"]], [["订单号"], ["=1+1"]]]):
            with self.subTest(index=index):
                source = self.workbook(f"cli-{index}.xlsx", {"订单": rows})
                manifest_path = self.root / f"manifest-{index}.json"
                manifest_path.write_text(json.dumps(self.manifest(source), ensure_ascii=False), encoding="utf-8")
                output = self.root / f"output-{index}"
                result = subprocess.run([sys.executable, str(script), "--manifest", str(manifest_path), "--output-dir", str(output)],
                                        capture_output=True, text=True, encoding="utf-8", timeout=10)
                events = [json.loads(line) for line in result.stdout.splitlines()]
                self.assertTrue(events)
                self.assertTrue(all(event["type"] == "progress" for event in events))
                if index == 0:
                    self.assertEqual(result.returncode, 0, result.stderr)
                    self.assertEqual(result.stderr, "")
                    self.assertTrue((output / "merged-001.xlsx").exists())
                else:
                    self.assertNotEqual(result.returncode, 0)
                    self.assertIn("合并失败：发现公式单元格", result.stderr)
                    self.assert_no_results(output)


if __name__ == "__main__":
    unittest.main()
