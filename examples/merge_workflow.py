"""Confirmed field rules, per-file checkpoints, and complete input accounting."""
import hashlib
import json
from pathlib import Path
import shutil
import tempfile

from openpyxl import load_workbook

from merge_adapter import SOURCE_COLUMNS, EXCEL_MAX_ROWS, _SplitWriter, _empty

PIPELINE_VERSION = 1


class InputError(ValueError):
    def __init__(self, message, sheet=None, row=None):
        super().__init__(message)
        self.sheet, self.row = sheet, row


def _header(values, sheet, row):
    values = list(values)
    while values and _empty(values[-1]):
        values.pop()
    if not values or any(not isinstance(value, str) or not value.strip() for value in values):
        raise InputError("表头必须是非空文字字段", sheet, row)
    if len(values) > 512 or len(set(values)) != len(values):
        raise InputError("表头重复或字段数超过 512", sheet, row)
    if any(value in SOURCE_COLUMNS for value in values):
        raise InputError("表头与保留来源字段重名", sheet, row)
    return values


def discover(entry):
    workbook = load_workbook(entry["path"], read_only=True, data_only=False)
    sheets = []
    try:
        for worksheet in workbook.worksheets:
            worksheet.reset_dimensions()
            rows = worksheet.iter_rows()
            for number, cells in enumerate(rows, 1):
                values = [cell.value for cell in cells]
                if all(_empty(value) for value in values):
                    continue
                try:
                    headers = _header(values, worksheet.title, number)
                except InputError as error:
                    sheets.append({"name": worksheet.title, "headerRow": number, "headers": [], "samples": [], "error": str(error)})
                    break
                samples = []
                for _ in range(3):
                    sample = next(rows, None)
                    if sample is None:
                        break
                    samples.append([str(cell.value)[:200] if cell.value is not None else None for cell in sample[:20]])
                sheets.append({"name": worksheet.title, "headerRow": number, "headers": headers, "samples": samples})
                break
    finally:
        workbook.close()
    return sheets


def _hash(path):
    digest = hashlib.sha256()
    with open(path, "rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _mapping(headers, rules, columns, sheet, row):
    configured = next((item for item in rules.get("mappings", []) if set(item["headers"]) == set(headers)), None)
    targets = configured["targets"] if configured else {name: name for name in headers}
    selected = {target: index for index, source in enumerate(headers) if (target := targets.get(source)) is not None}
    if len(selected) != len([value for value in targets.values() if value is not None]):
        raise InputError("字段映射存在重复目标", sheet, row)
    missing, extra = set(columns) - set(selected), set(selected) - set(columns)
    if missing or extra:
        raise InputError(f"字段不一致：缺少 {', '.join(sorted(missing)) or '无'}；新增 {', '.join(sorted(extra)) or '无'}", sheet, row)
    return [selected[name] for name in columns]


def _prepare_one(entry, rules, columns, cache_dir):
    actual_hash = _hash(entry["path"])
    if entry.get("sha256") and actual_hash != entry["sha256"]:
        raise InputError("输入内容与上传校验值不一致，请替换并重新上传")
    source = entry.get("relativePath") or entry["name"]
    key = hashlib.sha256(json.dumps([PIPELINE_VERSION, actual_hash, rules, columns, source], ensure_ascii=False, sort_keys=True).encode()).hexdigest()
    target = cache_dir / key
    try:
        metadata = json.loads((target / "metadata.json").read_text(encoding="utf-8"))
        if isinstance(metadata, dict) and isinstance(metadata.get("outputs"), list) and isinstance(metadata.get("hashes"), dict) and metadata["outputs"] and all(isinstance(name, str) and Path(name).name == name and (target / name).is_file() and _hash(target / name) == metadata["hashes"].get(name) for name in metadata["outputs"]):
            return {**metadata, "cacheKey": key, "cached": True}
    except (OSError, ValueError, KeyError):
        pass
    staging = Path(tempfile.mkdtemp(prefix=".prepare-", dir=cache_dir))
    writer = _SplitWriter(staging, columns + list(SOURCE_COLUMNS), EXCEL_MAX_ROWS)
    record_count, sheet_count = 0, 0
    workbook = None
    try:
        workbook = load_workbook(entry["path"], read_only=True, data_only=False)
        for worksheet in workbook.worksheets:
            if rules.get("sheetNames") and worksheet.title not in rules["sheetNames"]:
                continue
            worksheet.reset_dimensions()
            indexes, header_size = None, 0
            for row_number, cells in enumerate(worksheet.iter_rows(), 1):
                if any(cell.data_type == "f" for cell in cells):
                    raise InputError("发现公式单元格，请提供值数据", worksheet.title, row_number)
                values = [cell.value for cell in cells]
                if all(_empty(value) for value in values):
                    continue
                if indexes is None:
                    headers = _header(values, worksheet.title, row_number)
                    indexes = _mapping(headers, rules, columns, worksheet.title, row_number)
                    header_size = len(headers)
                    sheet_count += 1
                    continue
                if any(not _empty(value) for value in values[header_size:]):
                    raise InputError("存在没有对应表头的数据列", worksheet.title, row_number)
                values.extend([None] * max(0, header_size - len(values)))
                writer.append([values[index] for index in indexes] + [source, worksheet.title, row_number])
                record_count += 1
        if not sheet_count:
            raise InputError("没有选中的有效工作表，请检查工作表选择")
        writer.finish()
        metadata = {"recordCount": record_count, "sheetCount": sheet_count, "outputs": writer.outputs, "hashes": {name: _hash(staging / name) for name in writer.outputs}}
        (staging / "metadata.json").write_text(json.dumps(metadata, ensure_ascii=False), encoding="utf-8")
        if target.exists():
            assert target.parent == cache_dir
            shutil.rmtree(target)
        staging.rename(target)
        staging = None
        return {**metadata, "cacheKey": key, "cached": False}
    finally:
        if workbook is not None:
            workbook.close()
        if staging is not None:
            try:
                writer.discard()
            finally:
                shutil.rmtree(staging)


def prepare_files(manifest, on_progress=None):
    workflow = manifest["workflow"]
    rules = workflow["rules"]
    cache_dir = Path(workflow["cacheDir"]).resolve()
    cache_dir.mkdir(parents=True, exist_ok=True)
    files, issues, schemas = [], [], {}
    columns = list(rules.get("columns", []))
    for index, entry in enumerate(manifest["files"]):
        item = {"id": entry["id"], "relativePath": entry["relativePath"], "sheets": [], "status": "failed"}
        try:
            item["sheets"] = discover(entry)
            for sheet in item["sheets"]:
                if not sheet["headers"]:
                    continue
                signature = json.dumps(sorted(sheet["headers"]), ensure_ascii=False)
                schema = schemas.setdefault(signature, {"headers": sheet["headers"], "count": 0, "sample": sheet["samples"]})
                schema["count"] += 1
            selected = [sheet for sheet in item["sheets"] if sheet["headers"] and (not rules.get("sheetNames") or sheet["name"] in rules["sheetNames"])]
            if not columns and selected:
                columns = selected[0]["headers"]
            item.update(_prepare_one(entry, rules, columns, cache_dir))
            item["status"] = "success"
        except Exception as error:
            issue = {"fileId": entry["id"], "relativePath": entry["relativePath"], "sheet": getattr(error, "sheet", None), "row": getattr(error, "row", None), "message": str(error)[:1500]}
            issues.append(issue)
            item["error"] = issue
        files.append(item)
        if on_progress:
            on_progress({"type": "progress", "phase": "checking", "current": index + 1, "total": len(manifest["files"]), "message": f"已检查 {index + 1} / {len(manifest['files'])}：{entry['relativePath']}"})
    hashes = {}
    for entry in manifest["files"]:
        if entry.get("sha256"):
            hashes.setdefault(entry["sha256"], []).append(entry["relativePath"])
    return {"ok": bool(files) and not issues, "inputRevision": workflow["inputRevision"], "columns": columns, "files": files, "schemas": list(schemas.values()), "issues": issues,
            "fileCount": len(files), "successCount": sum(item["status"] == "success" for item in files), "failedCount": len(issues),
            "recordCount": sum(item.get("recordCount", 0) for item in files), "sheetCount": sum(item.get("sheetCount", 0) for item in files),
            "cachedCount": sum(bool(item.get("cached")) for item in files), "excluded": manifest.get("excluded", []),
            "duplicates": [names for names in hashes.values() if len(names) > 1]}


def merge_prepared(manifest, output_dir, on_progress=None):
    report = prepare_files(manifest, on_progress)
    if not report["ok"]:
        raise ValueError(f"{report['failedCount']} 份输入未通过检查，请修正异常后局部重试")
    output_dir = Path(output_dir).resolve()
    if output_dir.exists() and any(output_dir.iterdir()):
        raise ValueError("输出目录必须为空")
    staging = Path(tempfile.mkdtemp(prefix=".merge-", dir=output_dir.parent))
    writer = _SplitWriter(staging, report["columns"] + list(SOURCE_COLUMNS), EXCEL_MAX_ROWS)
    try:
        cache_dir = Path(manifest["workflow"]["cacheDir"])
        for index, item in enumerate(report["files"], 1):
            for name in item["outputs"]:
                workbook = load_workbook(cache_dir / item["cacheKey"] / name, read_only=True, data_only=False)
                try:
                    rows = workbook.active.iter_rows(values_only=True)
                    next(rows)
                    for row in rows:
                        writer.append(list(row))
                finally:
                    workbook.close()
            if on_progress:
                on_progress({"type": "progress", "phase": "merging", "current": index, "total": len(report["files"]), "message": f"已汇总 {index} 份已检查输入"})
        if on_progress:
            on_progress({"type": "progress", "phase": "exporting", "message": "正在保存并等待成果验收"})
        writer.finish()
        summary = {key: report[key] for key in ["fileCount", "sheetCount", "recordCount", "cachedCount", "excluded", "duplicates"]}
        summary.update({"outputCount": len(writer.outputs), "description": "按确认字段映射合并；无业务去重", "rules": manifest["workflow"]["rules"], "ruleVersion": PIPELINE_VERSION, "template": manifest["workflow"].get("template")})
        (staging / "summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
        if output_dir.exists():
            output_dir.rmdir()
        staging.rename(output_dir)
        staging = None
        return summary
    finally:
        if staging is not None:
            try:
                writer.discard()
            finally:
                shutil.rmtree(staging)


def preview_workbook(path, sheet=None, offset=0, limit=25):
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        names = workbook.sheetnames
        if sheet and sheet not in names:
            raise ValueError("工作表不存在")
        worksheet = workbook[sheet or names[0]]
        worksheet.reset_dimensions()
        rows = worksheet.iter_rows(values_only=True)
        header = next(rows, ())
        if len(header) > 512:
            raise ValueError("预览最多支持 512 列")
        result, more = [], False
        for index, row in enumerate(rows):
            if index < offset:
                continue
            if len(result) >= limit:
                more = True
                break
            result.append([str(value)[:2000] if value is not None else None for value in row])
        return {"sheetNames": names, "sheet": worksheet.title, "headers": [str(value) if value is not None else "" for value in header], "rows": result, "offset": offset, "hasMore": more}
    finally:
        workbook.close()


def validate_outputs(paths, expected_records=None, expected_headers=None):
    count, reports = 0, []
    for path in paths:
        workbook = load_workbook(path, read_only=True, data_only=False)
        file_count = 0
        try:
            for worksheet in workbook.worksheets:
                worksheet.reset_dimensions()
                rows = worksheet.iter_rows(values_only=True)
                header = next(rows, None)
                if header is None or all(_empty(value) for value in header):
                    raise ValueError(f"{Path(path).name} / {worksheet.title} 缺少结果表头")
                if expected_headers is not None and list(header) != expected_headers:
                    raise ValueError(f"{Path(path).name} / {worksheet.title} 成果字段与已确认规则不一致")
                file_count += sum(1 for row in rows if any(not _empty(value) for value in row))
        finally:
            workbook.close()
        reports.append({"name": Path(path).name, "recordCount": file_count})
        count += file_count
    if expected_records is not None and count != expected_records:
        raise ValueError(f"成果记录数不一致：预期 {expected_records}，实际 {count}")
    return {"ok": True, "recordCount": count, "files": reports}
